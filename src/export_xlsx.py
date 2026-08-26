"""Excel export matching docs/excel-columns.md.

Generates exports/job_outreach.xlsx with all 30 columns for every
researched lead, including WATCHLIST, REJECT, and blocked entries.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db

EXPORT_PATH = Path("exports/job_outreach.xlsx")

# Column definitions from docs/excel-columns.md
COLUMNS = [
    "Company",
    "Website",
    "Sector",
    "Employee Count",
    "Funding Status",
    "Funding Details",
    "Founder(s)",
    "Founder University",
    "Founder Previous Companies",
    "Founder Quality Score",
    "Company Fit Score",
    "Company Classification",
    "Person",
    "Title",
    "Contact Priority",
    "Profile URL",
    "Honorific",
    "Honorific Confidence",
    "Email",
    "Email Source",
    "Email Verification Status",
    "Email Score",
    "Personalization Paragraph",
    "Personalization Sources",
    "Send Eligible",
    "Status",
    "Scheduled At",
    "Sent At",
    "Gmail Message ID",
    "Notes",
]


def export(db_path: str = str(db.DB_PATH), output_path: str = str(EXPORT_PATH)) -> Path:
    """Generate the Excel export with all leads and company data."""
    conn = db.get_connection(db_path)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Job Outreach"

    # ── Header row ────────────────────────────────────────────────────────
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", wrap_text=True)

    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # ── Data rows ─────────────────────────────────────────────────────────
    # Get all companies with their leads and founders
    companies = db.get_all_companies(conn)
    row_num = 2

    for company in companies:
        cid = company["id"]
        founders = db.get_founders_for_company(conn, cid)
        founder_names = ", ".join(f.get("full_name", "") for f in founders)
        founder_unis = ", ".join(f.get("university", "") or "" for f in founders)
        founder_prev = ", ".join(f.get("previous_companies", "") or "" for f in founders)

        # Get leads for this company
        leads = conn.execute(
            "SELECT * FROM leads WHERE company_id = ?", (cid,)
        ).fetchall()

        if not leads:
            # Company with no leads — still write one row
            ws.cell(row=row_num, column=1, value=company.get("name"))
            ws.cell(row=row_num, column=2, value=company.get("website"))
            ws.cell(row=row_num, column=3, value=company.get("sector"))
            ws.cell(row=row_num, column=4, value=company.get("employee_count"))
            ws.cell(row=row_num, column=5, value=company.get("funding_status"))
            ws.cell(row=row_num, column=6, value=company.get("funding_details"))
            ws.cell(row=row_num, column=7, value=founder_names)
            ws.cell(row=row_num, column=8, value=founder_unis)
            ws.cell(row=row_num, column=9, value=founder_prev)
            ws.cell(row=row_num, column=10, value=company.get("founder_quality_score"))
            ws.cell(row=row_num, column=11, value=company.get("total_fit_score"))
            ws.cell(row=row_num, column=12, value=company.get("classification"))
            ws.cell(row=row_num, column=30, value=company.get("rationale"))
            row_num += 1
            continue

        for lead in leads:
            lead = dict(lead)
            # Get send record if exists
            send = conn.execute(
                "SELECT * FROM sends WHERE lead_id = ? ORDER BY id DESC LIMIT 1",
                (lead["id"],)
            ).fetchone()
            send = dict(send) if send else {}

            ws.cell(row=row_num, column=1, value=company.get("name"))
            ws.cell(row=row_num, column=2, value=company.get("website"))
            ws.cell(row=row_num, column=3, value=company.get("sector"))
            ws.cell(row=row_num, column=4, value=company.get("employee_count"))
            ws.cell(row=row_num, column=5, value=company.get("funding_status"))
            ws.cell(row=row_num, column=6, value=company.get("funding_details"))
            ws.cell(row=row_num, column=7, value=founder_names)
            ws.cell(row=row_num, column=8, value=founder_unis)
            ws.cell(row=row_num, column=9, value=founder_prev)
            ws.cell(row=row_num, column=10, value=company.get("founder_quality_score"))
            ws.cell(row=row_num, column=11, value=company.get("total_fit_score"))
            ws.cell(row=row_num, column=12, value=company.get("classification"))
            ws.cell(row=row_num, column=13, value=lead.get("full_name"))
            ws.cell(row=row_num, column=14, value=lead.get("title"))
            ws.cell(row=row_num, column=15, value=lead.get("contact_priority"))
            ws.cell(row=row_num, column=16, value=lead.get("profile_url"))
            ws.cell(row=row_num, column=17, value=lead.get("honorific"))
            ws.cell(row=row_num, column=18, value=lead.get("honorific_confidence"))
            ws.cell(row=row_num, column=19, value=lead.get("email"))
            ws.cell(row=row_num, column=20, value=lead.get("email_provider"))
            ws.cell(row=row_num, column=21, value=lead.get("email_verification_status"))
            ws.cell(row=row_num, column=22, value=lead.get("email_score"))
            ws.cell(row=row_num, column=23, value=lead.get("personalization_paragraph"))
            ws.cell(row=row_num, column=24, value=lead.get("personalization_source_urls_json"))
            ws.cell(row=row_num, column=25, value="Yes" if lead.get("eligibility") else "No")
            ws.cell(row=row_num, column=26, value=lead.get("status"))
            ws.cell(row=row_num, column=27, value=send.get("scheduled_at"))
            ws.cell(row=row_num, column=28, value=send.get("sent_at"))
            ws.cell(row=row_num, column=29, value=send.get("gmail_message_id"))
            ws.cell(row=row_num, column=30, value=company.get("rationale"))
            row_num += 1

    # ── Auto-size columns ─────────────────────────────────────────────────
    for col_idx in range(1, len(COLUMNS) + 1):
        max_len = len(COLUMNS[col_idx - 1])
        for row in range(2, min(row_num, 20)):
            val = ws.cell(row=row, column=col_idx).value
            if val:
                max_len = max(max_len, min(len(str(val)), 50))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 3

    # ── Color code classifications ────────────────────────────────────────
    outreach_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    watchlist_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    reject_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")

    for row in range(2, row_num):
        classification = ws.cell(row=row, column=12).value
        if classification == "OUTREACH":
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(row=row, column=col).fill = outreach_fill
        elif classification == "WATCHLIST":
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(row=row, column=col).fill = watchlist_fill
        elif classification == "REJECT":
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(row=row, column=col).fill = reject_fill

    wb.save(str(output))
    conn.close()
    return output
